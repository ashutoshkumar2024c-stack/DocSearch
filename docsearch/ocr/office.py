"""
Native parsers for non-PDF formats, producing layout elements.

    .docx  body in reading order: headings (Heading/Title styles), paragraphs, tables;
           embedded pictures -> figure tool
    .pptx  per slide: title -> heading, text boxes, tables, pictures -> figure tool, speaker notes
    .xlsx  per sheet: a table element (header row + data rows)
    .csv   one table element
    .html  h1-h6 -> headings, paragraphs / list items, tables; scripts and styles removed
    text   paragraphs

`figure_fn(image, label) -> Element | None` is supplied by the Extractor so that
embedded images go through the same automatic OCR / vision routing as scans.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Callable

from PIL import Image

from .layout import is_real_table, table_to_markdown, text_to_elements
from .types import Element, PagePlan, PageText

FigureFn = Callable[[Image.Image, str], "Element | None"]
MIN_FIGURE_PX = 150


def _open_image(blob: bytes) -> Image.Image | None:
    try:
        img = Image.open(io.BytesIO(blob))
        img.load()
    except Exception:
        return None
    if min(img.size) < MIN_FIGURE_PX or max(img.size) / max(1, min(img.size)) > 12:
        return None                                      # icons, bullets, lines
    return img.convert("RGB")


# ------------------------------------------------------------------ DOCX ---
def parse_docx(path: Path, source: str, figure_fn: FigureFn | None) -> list[PageText]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    elements: list[Element] = []
    figures = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            style = (p.style.name if p.style is not None else "") or ""
            if text:
                m = re.match(r"Heading (\d)", style)
                if style == "Title":
                    elements.append(Element("heading", text, 1, "python-docx"))
                elif m:
                    elements.append(Element("heading", text, int(m.group(1)), "python-docx"))
                else:
                    elements.append(Element("text", text, tool="python-docx"))
            if figure_fn:
                for blip in child.iter(qn("a:blip")):
                    rid = blip.get(qn("r:embed"))
                    part = doc.part.related_parts.get(rid) if rid else None
                    img = _open_image(part.blob) if part is not None else None
                    if img is not None:
                        figures += 1
                        el = figure_fn(img, f"figure {figures}")
                        if el:
                            elements.append(el)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            rows = []
            for row in table.rows:
                cells, prev = [], None
                for c in row.cells:                      # merged cells repeat; keep once
                    if c._tc is not prev:
                        cells.append(c.text)
                    prev = c._tc
                rows.append(cells)
            if is_real_table(rows):
                elements.append(Element("table", table_to_markdown(rows), tool="python-docx"))
            else:
                flat = " ".join(" ".join(r) for r in rows).strip()
                if flat:
                    elements.append(Element("text", flat, tool="python-docx"))
    plan = PagePlan(None, "docx", "parser", ["python-docx"] + (["figure-router"] if figures else []),
                    [f"Word document: parsed natively ({figures} embedded image(s) sent to figure routing)"])
    return [PageText.from_elements(source, elements, method="parser", route="parser",
                                   content_type="docx", plan=plan)]


# ------------------------------------------------------------------ PPTX ---
def _max_font_pt(text_frame) -> float:
    sizes = [r.font.size.pt for p in text_frame.paragraphs for r in p.runs if r.font.size]
    return max(sizes) if sizes else 0.0


def parse_pptx(path: Path, source: str, figure_fn: FigureFn | None) -> list[PageText]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    pages = []
    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        elements: list[Element] = []
        figures = 0
        title_shape = slide.shapes.title

        def walk(shapes):
            nonlocal figures
            for shape in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes)
                    continue
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[c.text for c in r.cells] for r in shape.table.rows]
                    if is_real_table(rows):
                        elements.append(Element("table", table_to_markdown(rows), tool="python-pptx"))
                elif getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    elements.append(Element("text", shape.text_frame.text.strip(), tool="python-pptx"))
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and figure_fn:
                    img = _open_image(shape.image.blob)
                    if img is not None:
                        figures += 1
                        el = figure_fn(img, f"slide {n} picture {figures}")
                        if el:
                            elements.append(el)

        if title_shape is not None and title_shape.has_text_frame and title_shape.text_frame.text.strip():
            elements.append(Element("heading", title_shape.text_frame.text.strip(), 1, "python-pptx"))
        walk(slide.shapes)
        if title_shape is None or not title_shape.text_frame.text.strip():
            # decks built from plain text boxes: the largest short text is the slide title
            sized = [(s, _max_font_pt(s.text_frame)) for s in slide.shapes
                     if getattr(s, "has_text_frame", False) and s.text_frame.text.strip()]
            sized = [(s, pt) for s, pt in sized if pt >= 20 and len(s.text_frame.text.strip()) <= 100
                     and any(c.isalpha() for c in s.text_frame.text)]
            if sized:
                best = max(sized, key=lambda x: x[1])[0].text_frame.text.strip()
                for i, e in enumerate(elements):
                    if e.kind == "text" and e.text == best:
                        elements.insert(0, Element("heading", elements.pop(i).text, 1, "python-pptx"))
                        break
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            elements.append(Element("text", "Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip(),
                                    tool="python-pptx"))
        if elements:
            plan = PagePlan(n, "pptx", "parser", ["python-pptx"] + (["figure-router"] if figures else []),
                            [f"slide parsed natively ({figures} picture(s) sent to figure routing)"])
            pages.append(PageText.from_elements(source, elements, page=n, method="parser", route="parser",
                                                content_type="presentation", plan=plan))
    return pages


# ---------------------------------------------------------- spreadsheets ---
def parse_xlsx(path: Path, source: str) -> list[PageText]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    pages = []
    try:
        for n, sheet in enumerate(wb.worksheets, 1):
            rows = [["" if v is None else v for v in r] for r in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if any(str(v).strip() for v in r)]
            if not rows:
                continue
            el = [Element("heading", f"Worksheet: {sheet.title}", 1, "openpyxl"),
                  Element("table", table_to_markdown(rows), tool="openpyxl")]
            plan = PagePlan(n, "xlsx", "parser", ["openpyxl"], [f"spreadsheet sheet '{sheet.title}': {len(rows)} rows"])
            pages.append(PageText.from_elements(source, el, page=n, method="parser", route="parser",
                                                content_type="spreadsheet", plan=plan))
    finally:
        wb.close()
    return pages


def parse_csv(path: Path, source: str, text: str) -> list[PageText]:
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    plan = PagePlan(None, "csv", "parser", ["csv"], [f"CSV table: {len(rows)} rows"])
    return [PageText.from_elements(source, [Element("table", table_to_markdown(rows), tool="csv")],
                                   method="parser", route="parser", content_type="spreadsheet", plan=plan)]


# ------------------------------------------------------------------ HTML ---
def parse_html(path: Path, source: str, text: str) -> list[PageText]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(text, "html.parser")
    for el in soup(["script", "style", "noscript", "template", "svg", "nav", "footer"]):
        el.decompose()
    elements: list[Element] = []
    if soup.title and soup.title.get_text(strip=True):
        elements.append(Element("heading", soup.title.get_text(" ", strip=True), 1, "beautifulsoup"))
    body = soup.body or soup
    for node in body.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table", "pre", "blockquote"]):
        if node.find_parent("table") is not None and node.name != "table":
            continue
        if node.name == "table":
            rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])] for tr in node.find_all("tr")]
            if is_real_table(rows):
                elements.append(Element("table", table_to_markdown(rows), tool="beautifulsoup"))
            continue
        t = node.get_text(" ", strip=True)
        if not t or (node.name == "p" and node.find_parent("li")):
            continue
        if node.name[0] == "h":
            elements.append(Element("heading", t, int(node.name[1]), "beautifulsoup"))
        else:
            elements.append(Element("text", ("• " if node.name == "li" else "") + t, tool="beautifulsoup"))
    if len(elements) <= 1:                     # pages built from bare divs
        elements += text_to_elements(body.get_text("\n", strip=True), "beautifulsoup")
    plan = PagePlan(None, "html", "parser", ["beautifulsoup"], ["web page: visible text, headings and tables"])
    return [PageText.from_elements(source, elements, method="parser", route="parser",
                                   content_type="html", plan=plan)]


def parse_text(path: Path, source: str, text: str) -> list[PageText]:
    kind = "markdown" if path.suffix.lower() == ".md" else "text"
    plan = PagePlan(None, kind, "parser", ["text"], [f"{kind} file"])
    return [PageText.from_elements(source, text_to_elements(text, "text"), method="parser",
                                   route="parser", content_type="text", plan=plan)]
