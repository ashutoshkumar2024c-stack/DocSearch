"""Build the manual test pack:  python tests/make_testpack.py  ->  test_documents/
One file per document type, each with known facts (see test_documents/QUESTIONS.md)."""

from __future__ import annotations

import csv
import random
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from make_samples import font, scanned_page   # noqa: E402,F401

from PIL import Image, ImageDraw, ImageFilter  # noqa: E402

OUT = Path(__file__).resolve().parent.parent / "test_documents"
OUT.mkdir(exist_ok=True)


# 1. digital PDF: headings, paragraphs, a table ------------------------------------------------
def leave_policy():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(OUT / "01_leave_policy_digital.pdf"), pagesize=A4)
    t = Table([["Leave Type", "Days per Year", "Notes"],
               ["Annual leave", "24", "Paid; up to 10 days may be carried forward"],
               ["Sick leave", "12", "Medical certificate needed after 3 days"],
               ["Casual leave", "8", "Cannot be combined with annual leave"],
               ["Maternity leave", "182", "Applies to the first two children"],
               ["Paternity leave", "15", "Must be taken within 6 months of birth"]])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.6, colors.black),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    doc.build([
        Paragraph("Greenfield Institute Employee Leave Policy", st["Title"]),
        Paragraph("1. Eligibility", st["Heading1"]),
        Paragraph("All full-time employees who have completed their probation period of 6 months are eligible "
                  "for paid leave. Part-time staff receive leave in proportion to their working hours.", st["BodyText"]),
        Paragraph("2. Leave entitlement", st["Heading1"]),
        Paragraph("The entitlement for each leave type is shown in the table below.", st["BodyText"]),
        Spacer(1, 8), t, Spacer(1, 12),
        Paragraph("3. Applying for leave", st["Heading1"]),
        Paragraph("Leave must be requested through the HR portal at least 5 working days in advance. "
                  "Emergency leave can be reported to the department head by phone.", st["BodyText"]),
    ])


# 2. scanned PDF: invoice with a ruled table (image only, no text layer) -------------------------
def invoice_scan():
    random.seed(3)
    w, h = 1654, 2339
    img = Image.new("L", (w, h), 244)
    d = ImageDraw.Draw(img)
    d.text((150, 160), "TAX INVOICE", fill=20, font=font(54))
    for i, line in enumerate(["Invoice No: INV-2026-0417", "Date: 03/10/2026",
                              "Bill To: Sunrise Pharma Pvt Ltd", "Payment terms: net 30 days"]):
        d.text((150, 280 + i * 62), line, fill=25, font=font(34))
    rows = [["Item", "Qty", "Rate (Rs)", "Amount (Rs)"], ["Laptop stand", "5", "1,200", "6,000"],
            ["USB hub", "10", "650", "6,500"], ["Wireless mouse", "8", "450", "3,600"]]
    x0, y0, cw, rh = 150, 640, [560, 160, 300, 380], 80
    for r, row in enumerate(rows):
        x = x0
        for c, cell in enumerate(row):
            d.rectangle([x, y0 + r * rh, x + cw[c], y0 + (r + 1) * rh], outline=30, width=3)
            d.text((x + 16, y0 + r * rh + 20), cell, fill=20, font=font(34))
            x += cw[c]
    d.text((150, y0 + 4 * rh + 60), "Sub-total Rs 16,100   GST 18%: Rs 2,898   Grand total: Rs 18,998",
           fill=20, font=font(34))
    img = img.rotate(-1.2, resample=Image.BICUBIC, fillcolor=244)
    px = img.load()
    for _ in range(8000):
        px[random.randrange(w), random.randrange(h)] = random.choice((45, 205))
    img = img.filter(ImageFilter.GaussianBlur(0.6)).convert("RGB")
    img.save(OUT / "02_invoice_scanned.pdf", resolution=200)


# 3. photographed receipt: rotated, dim, uneven light, noise, JPEG ------------------------------
def receipt_photo():
    random.seed(5)
    w, h = 900, 1200
    img = Image.new("L", (w, h), 225)
    d = ImageDraw.Draw(img)
    lines = ["CAFE BRAVO", "12 Lake Road, Pune", "Date: 14/10/2026", "",
             "Cappuccino        180", "Veg Sandwich      240", "Chocolate Brownie 120", "",
             "TOTAL             540", "Paid by UPI. Thank you!"]
    y = 90
    for i, ln in enumerate(lines):
        d.text((90, y), ln, fill=35, font=font(44 if i == 0 else 36))
        y += 78 if i == 0 else 62
    img = img.rotate(6, resample=Image.BICUBIC, fillcolor=150, expand=True)
    w, h = img.size
    px = img.load()
    for yy in range(h):                         # uneven lighting: darker toward one corner
        for xx in range(0, w):
            px[xx, yy] = max(0, int(px[xx, yy] * (0.62 + 0.38 * (1 - (xx / w * 0.6 + yy / h * 0.4)))))
    for _ in range(14000):
        px[random.randrange(w), random.randrange(h)] = random.choice((40, 190))
    img = img.filter(ImageFilter.GaussianBlur(1.0)).convert("RGB")
    img.save(OUT / "03_receipt_photo.jpg", quality=55)


# 4. Word document -------------------------------------------------------------------------------
def minutes_docx():
    from docx import Document
    d = Document()
    d.add_heading("Tech Fest Committee: Meeting Minutes", 0)
    d.add_paragraph("Date: 9 October 2026. Attendees: Meera, Arjun, Neha and Prof. Rao.")
    d.add_heading("Decisions", 1)
    d.add_paragraph("The annual tech fest will be held on 12 December 2026 in the main auditorium.")
    d.add_paragraph("The total budget approved for the event is Rs 4,50,000.")
    d.add_heading("Action items", 1)
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, ["Action", "Owner", "Due date"]):
        c.text = txt
    for row in [("Finalise vendor shortlist", "Meera", "28 October"),
                ("Prepare budget draft", "Arjun", "2 November"),
                ("Book the auditorium", "Neha", "20 October")]:
        for c, txt in zip(t.add_row().cells, row):
            c.text = txt
    d.save(OUT / "04_meeting_minutes.docx")


# 5. PowerPoint ----------------------------------------------------------------------------------
def launch_pptx():
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text, s.placeholders[1].text = "Project Aurora Launch Plan", "Product team, Q4 2026"
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Goals"
    s.placeholders[1].text_frame.text = "Reach 10,000 active users within 6 months"
    for t in ["Reduce onboarding time to under 5 minutes", "Achieve a 4.5 star app store rating"]:
        s.placeholders[1].text_frame.add_paragraph().text = t
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Budget"
    tbl = s.shapes.add_table(4, 2, 914400, 1800000, 7000000, 1800000).table
    for r, row in enumerate([("Phase", "Cost (Rs lakh)"), ("Design", "12"), ("Development", "38"), ("Marketing", "20")]):
        for c, txt in enumerate(row):
            tbl.cell(r, c).text = txt
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Risks"
    s.placeholders[1].text_frame.text = "Third-party payment gateway approval may take 3 weeks"
    s.placeholders[1].text_frame.add_paragraph().text = "Two backend engineers are needed by November"
    p.save(OUT / "05_product_launch.pptx")


# 6. Excel ---------------------------------------------------------------------------------------
def sales_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Q3 Sales"
    ws.append(["Region", "Units sold", "Revenue (Rs)"])
    for r in [("North", 1200, 540000), ("South", 950, 427500), ("East", 780, 351000), ("West", 1430, 643500)]:
        ws.append(r)
    ws2 = wb.create_sheet("Targets")
    ws2.append(["Region", "Q4 target (Rs)"])
    for r in [("North", 600000), ("South", 480000), ("East", 400000), ("West", 700000)]:
        ws2.append(r)
    wb.save(OUT / "06_sales_data.xlsx")


# 7. CSV, 8. HTML, 9. Markdown -----------------------------------------------------------------
def small_text_files():
    with open(OUT / "07_students.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Roll No", "Name", "Department", "CGPA"])
        w.writerows([("21AI001", "Ananya Iyer", "AI & Data Engineering", 9.1), ("21AI002", "Rohan Mehta", "AI & Data Engineering", 8.4),
                     ("21CS014", "Kavya Nair", "Computer Science", 8.9), ("21EC033", "Imran Shaikh", "Electronics", 7.8),
                     ("21ME021", "Pooja Verma", "Mechanical", 8.2), ("21CE007", "Dev Patel", "Civil", 7.5)])
    (OUT / "08_library_faq.html").write_text("""<html><head><title>Library FAQ</title></head><body>
<h1>Central Library FAQ</h1>
<h2>Opening hours</h2><p>The library is open from 8 am to 10 pm on weekdays and from 9 am to 5 pm on Saturdays. It is closed on Sundays.</p>
<h2>Borrowing rules</h2><ul><li>Students can borrow up to 4 books at a time for 14 days.</li>
<li>Faculty can borrow up to 10 books for 60 days.</li><li>A late fee of Rs 5 per day applies to overdue books.</li></ul>
<h2>Lost books</h2><p>A lost book must be replaced or paid for at twice its listed price.</p></body></html>""", encoding="utf-8")
    (OUT / "09_lab_safety.md").write_text("""# Chemistry Lab Safety Guide

## Before you start
- Wear a lab coat, safety goggles and closed shoes at all times.
- Tie back long hair and remove loose jewellery.

## Handling chemicals
1. Always add acid to water, never water to acid.
2. Label every container with the name and date.
3. Dispose of solvents only in the red waste bin.

## Emergencies
In case of a spill on skin, rinse with running water for 15 minutes. The emergency contact number is 4455.
""", encoding="utf-8")


# 10. chart image --------------------------------------------------------------------------------
def chart_png():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    names, vals = ["CSE", "ECE", "Mech", "Civil", "Biotech"], [420, 310, 260, 180, 350]
    fig, ax = plt.subplots(figsize=(7, 4.2), dpi=120)
    bars = ax.bar(names, vals, color="#3b6fd4")
    ax.bar_label(bars, fontsize=12)
    ax.set_title("Enrollment by Department 2026", fontsize=15)
    ax.set_ylabel("Students")
    fig.tight_layout()
    fig.savefig(OUT / "10_enrollment_chart.png")


def main():
    leave_policy(); invoice_scan(); receipt_photo(); minutes_docx(); launch_pptx(); sales_xlsx()
    small_text_files(); chart_png()
    shutil.copy(Path(__file__).parent / "samples" / "mixed.pdf", OUT / "11_mixed_report.pdf")
    print("\n".join(sorted(p.name for p in OUT.iterdir())))


if __name__ == "__main__":
    main()
