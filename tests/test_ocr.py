"""Tests for classification, automatic tool routing, layout analysis and OCR.

Run:  python -m pytest -q
"""

from __future__ import annotations

import io
import json
import sys
import types
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter

from docsearch.ocr import (Extractor, ExtractionError, cer, choose_ocr, plan_pdf_page,
                           profile_image, summarize, wer)
from docsearch.ocr.engines import OCRResult, TesseractEngine, VisionLLMEngine
from docsearch.ocr.layout import table_to_markdown, text_to_elements
from docsearch.ocr.router import garbage_ratio

SAMPLES = Path(__file__).parent / "samples"
sys.path.insert(0, str(Path(__file__).parent))
import make_samples  # noqa: E402

CAPS_BASIC = {"tesseract": {"available": True, "languages": ["eng"]}, "easyocr": {"available": False},
              "paddleocr": {"available": False}, "vision": {"available": False},
              "pdfplumber": {"available": True}}


def caps(**on):
    c = json.loads(json.dumps(CAPS_BASIC))
    for k, v in on.items():
        c[k] = {"available": v} if isinstance(v, bool) else v
    return c


@pytest.fixture(scope="session", autouse=True)
def samples():
    if not (SAMPLES / "mixed.pdf").exists():
        make_samples.main()
    return json.loads((SAMPLES / "ground_truth.json").read_text())


# ------------------------------------------------------------ PDF routing ---
def test_router_decisions_on_mixed_pdf():
    plans = Extractor(caps=CAPS_BASIC).routes(SAMPLES / "mixed.pdf")
    assert [p.route for p in plans] == ["native", "ocr", "hybrid", "ocr"]
    assert plans[0].doc_type == "digital-pdf" and plans[1].doc_type == "scanned-pdf"


def test_pdf_page_rules():
    long = "This is a perfectly normal sentence with plenty of characters in it."
    assert plan_pdf_page(1, "", [], 600, 800, CAPS_BASIC).route == "ocr"
    assert plan_pdf_page(1, long, [], 600, 800, CAPS_BASIC).route == "native"
    assert plan_pdf_page(1, long, [(0, 0, 300, 400)], 600, 800, CAPS_BASIC).route == "hybrid"
    searchable = plan_pdf_page(1, long, [(0, 0, 600, 800)], 600, 800, CAPS_BASIC)
    assert searchable.route == "native" and searchable.doc_type == "searchable-scan"
    junk = "(cid:12)(cid:44)(cid:3)(cid:91)" * 10
    assert garbage_ratio(junk) > 0.3
    assert plan_pdf_page(1, junk, [], 600, 800, CAPS_BASIC).route == "ocr"


# -------------------------------------------------- image profile + tools ---
def test_image_profiles():
    scan = make_samples.scanned_page(make_samples.PAGE2)
    assert profile_image(scan).kind == "scan" and profile_image(scan).quality == "good"
    assert profile_image(make_samples.chart_image()).kind == "graphic"
    blurry = profile_image(scan.filter(ImageFilter.GaussianBlur(4)))
    assert blurry.quality == "poor" and "blurry" in blurry.issues
    faded = profile_image(ImageEnhance.Contrast(scan).enhance(0.25))
    assert "low contrast" in faded.issues
    assert profile_image(scan).script == "Latin"


def _profile(**kw):
    p = profile_image(make_samples.scanned_page(make_samples.PAGE2), detect_script=False)
    for k, v in kw.items():
        setattr(p, k, v)
    return p


def test_tool_assignment_matrix():
    clean, poor = _profile(), _profile(quality="poor", issues=("blurry",))
    graphic = _profile(kind="graphic")
    hindi = _profile(script="Devanagari")
    assert choose_ocr(clean, CAPS_BASIC)["engine"] == "tesseract"
    assert choose_ocr(clean, caps(vision=True))["fallback"] == "vision"
    assert choose_ocr(poor, CAPS_BASIC)["engine"] == "tesseract"            # nothing better installed
    assert choose_ocr(poor, caps(easyocr=True))["engine"] == "easyocr"
    assert choose_ocr(poor, caps(paddleocr=True))["engine"] == "paddleocr"
    g = choose_ocr(graphic, CAPS_BASIC)
    assert g["engine"] == "tesseract" and g["psm"] == 11
    g = choose_ocr(graphic, caps(vision=True))
    assert g["engine"] == "vision" and g["mode"] == "describe"
    hin = choose_ocr(hindi, caps(tesseract={"available": True, "languages": ["eng", "hin"]}))
    assert hin["engine"] == "tesseract" and hin["lang"] == "hin+eng"
    assert choose_ocr(hindi, caps(easyocr=True))["engine"] == "easyocr"
    assert choose_ocr(hindi, caps(vision=True))["engine"] == "vision"
    assert choose_ocr(clean, caps(easyocr=True), preferred="easyocr")["engine"] == "easyocr"
    # every decision explains itself
    assert all(choose_ocr(p, CAPS_BASIC)["reasons"] for p in (clean, poor, graphic, hindi))


# ------------------------------------------------------------- extraction ---
def test_mixed_pdf_extraction(samples):
    pages = Extractor(caps=CAPS_BASIC).extract(SAMPLES / "mixed.pdf")
    assert [p.page for p in pages] == [1, 2, 3, 4]
    p1, p2, p3, p4 = pages

    # page 1: digital text, heading and a real table
    assert cer("\n".join(samples["page1"][1:]), next(e.text for e in p1.elements if e.kind == "text")) == 0
    assert any(e.kind == "heading" and e.text == samples["page1"][0] for e in p1.elements)
    table = next(e.text for e in p1.elements if e.kind == "table")
    assert table.splitlines()[0] == "| " + " | ".join(samples["table1"][0]) + " |"
    assert "| Blue Ocean Ltd | 388 | 9 | 1,75,500 |" in table

    # page 2: scanned, noisy, rotated
    assert p2.route == "ocr" and p2.method == "tesseract" and p2.ocr_confidence > 80
    assert cer("\n".join(samples["page2"]), p2.text.replace("## ", "")) < 0.03

    # page 3: figure text paired with its values
    fig = next(e.text for e in p3.elements if e.kind == "figure")
    for item in samples["chart"]:
        assert item in fig

    # page 4: scanned ruled table read cell by cell
    t4 = next(e.text for e in p4.elements if e.kind == "table")
    assert "| Haemoglobin | 10.2 | g/dL |" in t4 and "| Glucose | 98 | mg/dL |" in t4
    assert any(e.kind == "heading" and e.text == samples["scan_table_title"] for e in p4.elements)

    s = summarize(pages)
    assert s["doc_type"] == "mixed-pdf" and s["elements"]["table"] == 2
    assert all(p.plan and p.plan.reasons for p in pages)


def test_skewed_photo(samples):
    (page,) = Extractor(caps=CAPS_BASIC).extract(SAMPLES / "photo.png")
    assert page.method == "tesseract" and page.content_type == "image"
    assert wer("\n".join(samples["photo"]), page.text.replace("## ", ""), lowercase=True) < 0.1


def test_upside_down_scan_is_turned_upright(tmp_path, samples):
    scan = make_samples.scanned_page(make_samples.PAGE2).rotate(180)
    scan.save(tmp_path / "flipped.png")
    (page,) = Extractor(caps=CAPS_BASIC).extract(tmp_path / "flipped.png")
    assert any("rotated 180" in r or "turned upright" in r for r in page.plan.reasons)
    assert "amoxicillin" in page.text


def test_office_formats_with_structure(tmp_path):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches, Pt

    chart = make_samples.chart_image()
    buf = io.BytesIO()
    chart.save(buf, format="PNG")

    d = Document()
    d.add_heading("Board Minutes", level=1)
    d.add_paragraph("The board approved the October budget.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Item", "Amount"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Budget", "4.2 crore"
    d.add_picture(io.BytesIO(buf.getvalue()), width=Inches(4))
    d.save(tmp_path / "a.docx")

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])                 # blank layout: no title placeholder
    tb = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1)).text_frame
    tb.text = "Release Roadmap"
    tb.paragraphs[0].runs[0].font.size = Pt(32)
    body = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text_frame
    body.text = "Ship the hybrid OCR router"
    body.paragraphs[0].runs[0].font.size = Pt(16)
    s.shapes.add_picture(io.BytesIO(buf.getvalue()), Inches(1), Inches(3), Inches(5))
    p.save(tmp_path / "b.pptx")

    (tmp_path / "c.html").write_text("<html><title>Policy</title><script>x=1</script><body>"
                                     "<h2>Refunds</h2><p>Refunds are issued within 14 days.</p>"
                                     "<table><tr><th>Plan</th><th>Days</th></tr><tr><td>Gold</td><td>30</td></tr></table>"
                                     "</body></html>")
    (tmp_path / "d.csv").write_text("name,score\nasha,91\nravi,84\n")

    ex = Extractor(caps=CAPS_BASIC)
    (docx,) = ex.extract(tmp_path / "a.docx")
    kinds = [e.kind for e in docx.elements]
    assert kinds[0] == "heading" and "table" in kinds and "figure" in kinds
    assert "| Budget | 4.2 crore |" in docx.text and "West 91%" in docx.text

    (slide,) = ex.extract(tmp_path / "b.pptx")
    assert slide.elements[0].kind == "heading" and slide.elements[0].text == "Release Roadmap"
    assert any(e.kind == "figure" and "North 82%" in e.text for e in slide.elements)

    (html,) = ex.extract(tmp_path / "c.html")
    assert "x=1" not in html.text and "| Gold | 30 |" in html.text
    assert any(e.kind == "heading" and e.text == "Refunds" for e in html.elements)
    (csvp,) = ex.extract(tmp_path / "d.csv")
    assert csvp.elements[0].kind == "table" and "| asha | 91 |" in csvp.text


def test_unreadable_files_fail_cleanly(tmp_path):
    Image.new("RGB", (400, 300), "white").save(tmp_path / "blank.png")
    with pytest.raises(ExtractionError):
        Extractor(caps=CAPS_BASIC).extract(tmp_path / "blank.png")
    (tmp_path / "x.bin").write_bytes(b"\x00\x01\x02" * 100)
    with pytest.raises(ExtractionError):
        Extractor(caps=CAPS_BASIC).extract(tmp_path / "x.bin")


def test_markdown_helpers():
    assert table_to_markdown([["a", "b"], ["1", None]]) == "| a | b |\n|---|---|\n| 1 |  |"
    els = text_to_elements("# Title\nintro line\n\n| x | y |\n|---|---|\n| 1 | 2 |", "vision")
    assert [e.kind for e in els] == ["heading", "text", "table"]


# ------------------------------------------------ engines, fallback, VLM ---
class _FakeOllama(types.ModuleType):
    def __init__(self, reply):
        super().__init__("ollama")
        self.reply, self.calls = reply, []
        outer = self

        class Client:
            def __init__(self, host=None):
                pass

            def chat(self, **kw):
                outer.calls.append(kw)
                return {"message": {"content": outer.reply}}

        self.Client = Client


def test_vision_engine_strips_thinking_and_sends_prompt(monkeypatch):
    fake = _FakeOllama("<think>let me look</think>\nINVOICE 42\nTotal: 900")
    monkeypatch.setitem(sys.modules, "ollama", fake)
    res = VisionLLMEngine(prompt="CUSTOM PROMPT").recognize(Image.new("RGB", (50, 50)))
    assert res.text == "INVOICE 42\nTotal: 900" and res.engine == "vision"
    user_msg = fake.calls[0]["messages"][1]
    assert "CUSTOM PROMPT" in user_msg["content"] and user_msg["images"]


def test_graphic_page_routed_to_vision_describe(monkeypatch, tmp_path):
    fake = _FakeOllama("Bar chart 'Renewal Rate by Region'.\n\n| Region | Rate |\n|---|---|\n| West | 91% |")
    monkeypatch.setitem(sys.modules, "ollama", fake)
    make_samples.chart_image().save(tmp_path / "chart.png")
    (page,) = Extractor(caps=caps(vision=True)).extract(tmp_path / "chart.png")
    assert page.route == "visual" and page.plan.tools == ["vision:describe"]
    assert page.elements[0].kind == "figure" and "West | 91%" in page.text
    assert "figure from a document" in fake.calls[0]["messages"][1]["content"]


def test_low_confidence_falls_back(tmp_path):
    class Weak:
        name = "tesseract"
        def recognize(self, image):
            return OCRResult("g4rbl3d t3xt h3re !!", confidence=31.0, engine="tesseract")

    class Strong:
        name = "vision"
        def recognize(self, image):
            return OCRResult("Handwritten note: call the supplier on Monday", engine="vision")

    Image.new("RGB", (300, 200), "white").save(tmp_path / "note.png")
    (page,) = Extractor(engine=Weak(), fallback_engine=Strong(), caps=CAPS_BASIC).extract(tmp_path / "note.png")
    assert page.method == "vision" and "supplier" in page.text
    assert any("retried with vision" in r for r in page.plan.reasons)


def test_failed_engine_degrades_to_tesseract(monkeypatch, tmp_path, samples):
    class Broken(types.ModuleType):
        def __init__(self):
            super().__init__("ollama")

            class Client:
                def __init__(self, host=None): pass
                def chat(self, **kw): raise ConnectionError("ollama is not running")
            self.Client = Client

    monkeypatch.setitem(sys.modules, "ollama", Broken())
    make_samples.scanned_page(make_samples.PAGE2).save(tmp_path / "scan.png")
    (page,) = Extractor(engine="vision", caps=caps(vision=True)).extract(tmp_path / "scan.png")
    assert page.method == "tesseract" and "amoxicillin" in page.text
    assert any("failed" in r for r in page.plan.reasons)


def test_easyocr_and_paddle_adapters(monkeypatch):
    easy = types.ModuleType("easyocr")

    class Reader:
        def __init__(self, langs, gpu=False, verbose=False): pass
        def readtext(self, arr, detail=1, paragraph=False):
            return [([[10, 10], [80, 10], [80, 30], [10, 30]], "Hello", 0.9),
                    ([[90, 12], [160, 12], [160, 30], [90, 30]], "world", 0.8),
                    ([[10, 80], [90, 80], [90, 100], [10, 100]], "Second", 0.7)]
    easy.Reader = Reader
    monkeypatch.setitem(sys.modules, "easyocr", easy)
    from docsearch.ocr.engines import EasyOCREngine, PaddleOCREngine, _easyocr_reader, _paddle
    _easyocr_reader.cache_clear()
    r = EasyOCREngine(gpu=False).recognize(Image.new("RGB", (200, 120)))
    assert r.text == "Hello world\n\nSecond" and 70 < r.confidence < 90

    paddle = types.ModuleType("paddleocr")

    class PaddleOCR:
        def __init__(self, **kw): pass
        def predict(self, arr):
            return [{"rec_texts": ["Total", "900"], "rec_scores": [0.99, 0.95],
                     "rec_polys": [[[0, 0], [50, 0], [50, 20], [0, 20]], [[60, 0], [90, 0], [90, 20], [60, 20]]]}]
    paddle.PaddleOCR = PaddleOCR
    monkeypatch.setitem(sys.modules, "paddleocr", paddle)
    _paddle.cache_clear()
    r = PaddleOCREngine().recognize(Image.new("RGB", (100, 30)))
    assert r.text == "Total 900" and r.confidence == pytest.approx(97.0)


def test_sparse_chart_ocr_pairs_labels_and_values():
    text = TesseractEngine().recognize_sparse(make_samples.chart_image()).text
    assert {"North 82%", "South 76%", "East 69%", "West 91%"} <= set(text.splitlines())
